import openpyxl
import os
import codecs
import sys

sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')

file_path = r"D:\Công việc\3. Dự án Sovigaz\2. Tài liệu phân tích và thiết kế\02. BRD & Template\Template báo cáo - BRD.xlsx"
output_dir = r"D:\Công việc\3. Dự án Sovigaz\AI-Document-Parser\output\documents"

wb = openpyxl.load_workbook(file_path, data_only=True)

target_sheets = [
    '1.1 temp doanh thu',
    '2.1 temp doanh thu chi tiết',
    '3.1 temp công nợ',
    '4.1 Temp KPI',
    '5.1 temp KPI chi tiết'
]

for sheet_name in target_sheets:
    if sheet_name not in wb.sheetnames:
        print(f"Bỏ qua: Không tìm thấy sheet '{sheet_name}'")
        continue
        
    print(f"Đang xử lý sheet: {sheet_name}")
    sheet = wb[sheet_name]
    
    # Chuẩn hóa tên file xuất ra (thay thế khoảng trắng bằng dấu gạch dưới)
    safe_name = sheet_name.replace(" ", "_").replace(".", "_") + ".md"
    output_path = os.path.join(output_dir, safe_name)
    
    md_lines = [f"# Giao diện mẫu (Template) - {sheet_name}\n"]
    
    for row in sheet.iter_rows(values_only=True):
        row_data = [str(cell).strip().replace("\n", " ") if cell is not None else "" for cell in row]
        
        # Trim empty columns at the end
        last_idx = -1
        for i in range(len(row_data)-1, -1, -1):
            if row_data[i] != "":
                last_idx = i
                break
                
        if last_idx == -1:
            md_lines.append("")
            continue
            
        cleaned_row = row_data[:last_idx+1]
        
        # Format thành dạng cột cách nhau bởi dấu |
        text = " **|** ".join([c for c in cleaned_row if c])
        md_lines.append(f"- {text}")
        
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
        
    print(f"-> Đã lưu thành công: {safe_name}")

print("HOÀN TẤT TOÀN BỘ!")
