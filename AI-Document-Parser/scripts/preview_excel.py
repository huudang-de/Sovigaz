import openpyxl
import sys
import codecs
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')

file_path = r"D:\Công việc\3. Dự án Sovigaz\2. Tài liệu phân tích và thiết kế\02. BRD & Template\Template báo cáo - BRD.xlsx"

try:
    print("Đang mở file Excel...")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    if "1.BRD- DOANH THU" in wb.sheetnames:
        sheet = wb["1.BRD- DOANH THU"]
        print(f"--- 15 DÒNG ĐẦU TIÊN CỦA SHEET '1.BRD- DOANH THU' ---")
        for row in sheet.iter_rows(min_row=1, max_row=15, values_only=True):
            print(row)
    else:
        print("Không tìm thấy Sheet '1.BRD- DOANH THU'. Các Sheet hiện có:")
        print(wb.sheetnames)
except Exception as e:
    print(f"Lỗi: {e}")
